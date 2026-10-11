from copy import deepcopy
from pathlib import Path
import shutil

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt


SOURCE = Path('/Users/skyler/Python Project/BlogSystem/软件综合实践报告_个人博客系统.docx')
OUTPUT = Path('/Users/skyler/Python Project/BlogSystem/软件综合实践报告_个人博客系统_按模板格式.docx')


def set_font_obj(font, name, size=None, bold=None):
    font.name = name
    if size is not None:
        font.size = Pt(size)
    if bold is not None:
        font.bold = bold


def set_rfonts(rpr, name):
    if rpr.rFonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.insert(0, rfonts)
    else:
        rfonts = rpr.rFonts
    for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        rfonts.set(qn(f'w:{key}'), name)


def format_run(run, name, size, bold=None):
    set_font_obj(run.font, name, size, bold)
    rpr = run._element.get_or_add_rPr()
    set_rfonts(rpr, name)


def format_style(style, name, size, bold=None):
    set_font_obj(style.font, name, size, bold)
    rpr = style._element.get_or_add_rPr()
    set_rfonts(rpr, name)


def format_body_paragraph(p):
    pf = p.paragraph_format
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.first_line_indent = Cm(0.74)
    pf.left_indent = None
    pf.right_indent = None
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = Pt(20)
    for run in p.runs:
        format_run(run, '宋体', 12)


def format_heading(p, level):
    pf = p.paragraph_format
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.first_line_indent = Pt(0)
    pf.left_indent = Pt(0)
    pf.right_indent = Pt(0)
    pf.space_before = Pt(6)
    pf.space_after = Pt(6)
    pf.line_spacing = Pt(20)
    pf.keep_with_next = True
    size = 16 if level == 1 else 14
    for run in p.runs:
        format_run(run, '黑体', size, False)


def format_list(p):
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = Pt(20)
    for run in p.runs:
        format_run(run, '宋体', 12)


def format_caption(p):
    pf = p.paragraph_format
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf.first_line_indent = Pt(0)
    pf.space_before = Pt(4)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.0
    pf.keep_with_next = True
    for run in p.runs:
        format_run(run, '宋体', 10.5, False)


def format_table(table):
    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            for p in cell.paragraphs:
                pf = p.paragraph_format
                pf.first_line_indent = Pt(0)
                pf.space_before = Pt(0)
                pf.space_after = Pt(0)
                pf.line_spacing = Pt(16)
                for run in p.runs:
                    # Keep existing bold header emphasis while normalizing type.
                    format_run(run, '宋体', 10.5)


def set_update_fields(doc):
    settings = doc.settings._element
    update = settings.find(qn('w:updateFields'))
    if update is None:
        update = OxmlElement('w:updateFields')
        settings.append(update)
    update.set(qn('w:val'), 'true')


shutil.copy2(SOURCE, OUTPUT)
doc = Document(OUTPUT)

# Match the reference page system while preserving section boundaries and chrome.
for idx, section in enumerate(doc.sections):
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    if idx == 0:
        section.left_margin = Inches(0.88)
        section.right_margin = Inches(0.64)
        section.top_margin = Inches(1.08)
        section.bottom_margin = Inches(1.00)
    else:
        section.left_margin = Inches(0.79)
        section.right_margin = Inches(0.79)
        section.top_margin = Inches(0.98)
        section.bottom_margin = Inches(0.98)

# Establish style defaults from the reference; direct formatting below makes the
# result deterministic even where the source has mixed run-level overrides.
style_specs = {
    'Report Normal': ('宋体', 12, None),
    'Report heading 1': ('黑体', 16, False),
    'Report heading 2': ('黑体', 14, False),
    'Report List Bullet': ('宋体', 12, None),
    'Report List Number': ('宋体', 12, None),
    'Report caption': ('宋体', 10.5, False),
    'toc 1': ('宋体', 12, None),
    'toc 2': ('宋体', 12, None),
}
for style_name, (font_name, size, bold) in style_specs.items():
    if style_name in [s.name for s in doc.styles]:
        format_style(doc.styles[style_name], font_name, size, bold)

for p in doc.paragraphs:
    style_name = p.style.name
    if style_name == 'Report heading 1':
        format_heading(p, 1)
    elif style_name == 'Report heading 2':
        format_heading(p, 2)
    elif style_name == 'Report Normal':
        # The only centered Report Normal paragraph is the TOC title.
        if p.text.strip().replace('\u3000', '').replace(' ', '') == '目录':
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf = p.paragraph_format
            pf.first_line_indent = Pt(0)
            pf.space_before = Pt(0)
            pf.space_after = Pt(0)
            pf.line_spacing = Pt(20)
            for run in p.runs:
                format_run(run, '黑体', 24, False)
        else:
            format_body_paragraph(p)
    elif style_name in ('Report List Bullet', 'Report List Number'):
        format_list(p)
    elif style_name == 'Report caption':
        format_caption(p)
    elif style_name in ('toc 1', 'toc 2'):
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = Pt(20)
        for run in p.runs:
            format_run(run, '宋体', 12)

# The first table is the template-derived cover name/student-number block.
for table in doc.tables[1:]:
    format_table(table)

set_update_fields(doc)
doc.save(OUTPUT)
print(OUTPUT)
