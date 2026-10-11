from pathlib import Path
from collections import Counter
from docx import Document
from docx.oxml.ns import qn


def fmt_run(run):
    rpr = run._element.rPr
    fonts = {}
    if rpr is not None and rpr.rFonts is not None:
        for key in ("ascii", "hAnsi", "eastAsia", "cs"):
            value = rpr.rFonts.get(qn(f"w:{key}"))
            if value:
                fonts[key] = value
    return {
        "text": run.text.replace("\t", "<TAB>").replace("\n", "<BR>"),
        "fonts": fonts,
        "font_name": run.font.name,
        "size_pt": run.font.size.pt if run.font.size else None,
        "bold": run.font.bold,
        "italic": run.font.italic,
        "underline": bool(run.font.underline) if run.font.underline is not None else None,
    }


def fmt_para(p):
    pf = p.paragraph_format
    return {
        "style": p.style.name,
        "alignment": str(p.alignment),
        "left_cm": pf.left_indent.cm if pf.left_indent else None,
        "right_cm": pf.right_indent.cm if pf.right_indent else None,
        "first_cm": pf.first_line_indent.cm if pf.first_line_indent else None,
        "before_pt": pf.space_before.pt if pf.space_before else None,
        "after_pt": pf.space_after.pt if pf.space_after else None,
        "line": str(pf.line_spacing),
        "keep_next": pf.keep_with_next,
        "page_break": pf.page_break_before,
        "text": p.text.replace("\t", "<TAB>").replace("\n", "<BR>"),
        "runs": [fmt_run(r) for r in p.runs if r.text or r._element.rPr is not None],
    }


def dump(path):
    doc = Document(path)
    print(f"\n===== {Path(path).name} =====")
    print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} sections={len(doc.sections)}")
    for i, p in enumerate(doc.paragraphs):
        info = fmt_para(p)
        if info["text"].strip() or i < 30:
            print(f"P{i}: {info}")
    for ti, table in enumerate(doc.tables):
        print(f"\nTABLE {ti}: rows={len(table.rows)} cols={len(table.columns)} style={table.style.name if table.style else None}")
        for ri, row in enumerate(table.rows):
            vals = []
            for ci, cell in enumerate(row.cells):
                vals.append({"text": cell.text.replace("\n", " / "), "paras": [fmt_para(p) for p in cell.paragraphs]})
            print(f"R{ri}: {vals}")
    print("\nSTYLES USED:")
    counts = Counter(p.style.name for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                counts.update(p.style.name for p in cell.paragraphs)
    for name, count in counts.most_common():
        st = doc.styles[name]
        print(name, count, {
            "font": st.font.name,
            "size": st.font.size.pt if st.font.size else None,
            "bold": st.font.bold,
            "eastAsia": st._element.rPr.rFonts.get(qn("w:eastAsia")) if st._element.rPr is not None and st._element.rPr.rFonts is not None else None,
            "alignment": str(st.paragraph_format.alignment),
            "line": str(st.paragraph_format.line_spacing),
            "before": st.paragraph_format.space_before.pt if st.paragraph_format.space_before else None,
            "after": st.paragraph_format.space_after.pt if st.paragraph_format.space_after else None,
        })


dump('/Users/skyler/Python Project/BlogSystem/.codex-doc-format-work/模板转换.docx')
dump('/Users/skyler/Python Project/BlogSystem/软件综合实践报告_个人博客系统.docx')
